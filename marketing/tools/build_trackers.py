#!/usr/bin/env python3
"""Génère les fichiers de suivi commerciaux (CSV FR/EN et classeur KPI avec formules).

  python3 marketing/tools/build_trackers.py
  python3 <skill xlsx>/scripts/recalc.py marketing/trackers/kpi-tableau-de-bord.xlsx

Les fichiers sont livrés vides (en-têtes seulement) : aucune donnée de prospect ou de
partenaire, réelle ou fictive. Le format de chaque colonne est décrit dans
marketing/fr/08-suivi-et-processus.md (en anglais : marketing/en/08-sales-process-and-tracking.md).
"""
import csv, os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.path.join(ROOT, "trackers")
os.makedirs(OUT, exist_ok=True)

PROSPECT_FR = ["id", "date_entree", "etablissement", "secteur", "commune", "contact_nom", "contact_role", "telephone",
               "email", "consentement_canal", "consentement_date", "source", "score", "statut", "derniere_action",
               "prochaine_action", "date_prochaine_action", "motif_perte", "notes"]
PROSPECT_EN = ["id", "entry_date", "business", "sector", "commune", "contact_name", "contact_role", "phone", "email",
               "consent_channel", "consent_date", "source", "score", "status", "last_action", "next_action",
               "next_action_date", "loss_reason", "notes"]
PARTNER_FR = ["id", "nom", "type", "commune", "date_contrat", "prospects_enregistres", "demos", "pilotes", "clients",
              "encaisse_usd", "commission_due_usd", "commission_payee_usd", "derniere_rencontre", "notes"]
PARTNER_EN = ["id", "name", "type", "commune", "contract_date", "prospects_recorded", "demos", "pilots", "customers",
              "collected_usd", "commission_due_usd", "commission_paid_usd", "last_meeting", "notes"]
STATUSES = ["identifié", "contacté", "démo fixée", "démo faite", "pilote proposé", "pilote en cours", "proposition",
            "client", "perdu", "plus tard"]

def write_csv(name, header):
    with open(os.path.join(OUT, name), "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f, delimiter=";").writerow(header)

write_csv("suivi-prospects.csv", PROSPECT_FR)
write_csv("prospect-tracker.csv", PROSPECT_EN)
write_csv("suivi-partenaires.csv", PARTNER_FR)
write_csv("partner-tracker.csv", PARTNER_EN)

# ------------------------------------------------------------------ workbook
F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="173F5F")
INP = Font(name=F, color="0000FF")
YEL = PatternFill("solid", fgColor="FFF2CC")
NORM = Font(name=F)
B = Font(name=F, bold=True)
thin = Side(style="thin", color="C9C2B3")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)

def header(ws, row, labels, widths=None):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.alignment, c.border = H, HF, Alignment(wrap_text=True, vertical="center"), BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

wb = Workbook()
# --- Taux / Rates
tx = wb.active; tx.title = "Taux"
tx["A1"] = "Taux de change de référence / Reference exchange rate"; tx["A1"].font = Font(name=F, bold=True, size=13)
rows = [("CDF pour 1 USD / CDF per 1 USD", 2264.0), ("Date du taux / Rate date", "2026-10-26"),
        ("Source", "Banque centrale du Congo, cours officiel — https://www.bcc.cd/marche-des-changes/cours-de-change")]
for r, (k, v) in enumerate(rows, 3):
    tx.cell(row=r, column=1, value=k).font = B
    c = tx.cell(row=r, column=2, value=v); c.font = INP; c.fill = YEL
tx["B3"].number_format = "#,##0.0000"
tx["B3"].comment = Comment("Taux officiel BCC du 26/10/2026 (consulté le 27/10/2026). Mettre à jour avant toute dépense.", "Nhill")
tx["A7"] = "Cellules jaunes à texte bleu = à saisir. Les autres cellules sont calculées. / Yellow cells with blue text = inputs; others are formulas."
tx["A7"].font = Font(name=F, italic=True)
tx.column_dimensions["A"].width = 38; tx.column_dimensions["B"].width = 90

# --- Budget
bd = wb.create_sheet("Budget")
bd["A1"] = "Budget 90 jours — enveloppes proposées (à confirmer par devis) / 90-day budget — proposed envelopes"
bd["A1"].font = Font(name=F, bold=True, size=13)
bd["A2"] = "Choix du scénario / Scenario (A ou/or B) :"; bd["A2"].font = B
bd["C2"] = "A"; bd["C2"].font = INP; bd["C2"].fill = YEL
dv = DataValidation(type="list", formula1='"A,B"', allow_blank=False); bd.add_data_validation(dv); dv.add("C2")
header(bd, 4, ["Poste / Item", "Enveloppe A (USD)", "Enveloppe B (USD)", "Enveloppe retenue (USD)", "Enveloppe retenue (CDF)",
               "Dépensé (USD)", "Dépensé (CDF)", "Écart (USD)", "Justificatif / Receipt"], [62, 16, 16, 18, 20, 14, 18, 14, 28])
items = [("Forfaits de données mobiles / Mobile data", 45, 60), ("Impression fiches et affiches / Sheets and posters", 60, 60),
         ("Impression complémentaire / Extra printing", 0, 150), ("Transport", 120, 250),
         ("Domaine et hébergement du site / Domain and hosting", 30, 30), ("Séance de démonstration groupée / Group demo", 40, 0),
         ("Publicité réseaux sociaux / Social-media ads", 0, 300), ("Événement d'entrepreneurs / Entrepreneurs' event", 0, 200),
         ("Formation des partenaires / Partner training", 0, 100), ("Micro et support vidéo / Mic and video stand", 0, 80),
         ("Tablette Android de démonstration / Android demo tablet", 0, 150)]
first = 5
for i, (name, a, b) in enumerate(items):
    r = first + i
    bd.cell(row=r, column=1, value=name).font = NORM
    for col, v in ((2, a), (3, b)):
        c = bd.cell(row=r, column=col, value=v); c.font = INP; c.number_format = "#,##0;(#,##0);-"
    bd.cell(row=r, column=4, value=f'=IF($C$2="B",C{r},B{r})').number_format = "#,##0;(#,##0);-"
    bd.cell(row=r, column=5, value=f"=D{r}*Taux!$B$3").number_format = "#,##0;(#,##0);-"
    c = bd.cell(row=r, column=6, value=0); c.font = INP; c.fill = YEL; c.number_format = "#,##0.00;(#,##0.00);-"
    bd.cell(row=r, column=7, value=f"=F{r}*Taux!$B$3").number_format = "#,##0;(#,##0);-"
    bd.cell(row=r, column=8, value=f"=D{r}-F{r}").number_format = "#,##0.00;(#,##0.00);-"
last = first + len(items) - 1
t = last + 1
bd.cell(row=t, column=1, value="Total").font = B
for col in range(2, 9):
    L = get_column_letter(col)
    c = bd.cell(row=t, column=col, value=f"=SUM({L}{first}:{L}{last})"); c.font = B
    c.number_format = "#,##0;(#,##0);-"
bd.cell(row=t + 2, column=1, value="Taux utilisé : feuille « Taux » (BCC). Enveloppes = propositions du plan de 90 jours, pas des prix relevés.").font = Font(name=F, italic=True)
BUDGET_TOTAL = f"Budget!$D${t}"; SPENT_TOTAL = f"Budget!$F${t}"

# --- Hebdo (weekly)
hb = wb.create_sheet("Hebdo")
cols = ["Semaine / Week", "Lundi / Monday", "Prospects ajoutés / added", "Contacts", "Démos faites / done",
        "Essais acceptés / trials accepted", "Pilotes actifs / active", "Pilotes terminés / completed", "Clients",
        "Partenaires signés / signed", "Note recommandation moy. / avg score", "Incidents bloquants > 48 h"]
header(hb, 1, cols, [10, 13, 13, 11, 12, 12, 12, 13, 10, 13, 16, 14])
for w in range(1, 14):
    r = w + 1
    hb.cell(row=r, column=1, value=f"S{w}").font = NORM
    c = hb.cell(row=r, column=2, value=f"=IF('Tableau de bord'!$B$3=\"\",\"\",'Tableau de bord'!$B$3+{(w-1)*7})"); c.number_format = "yyyy-mm-dd"
    for col in range(3, 13):
        cell = hb.cell(row=r, column=col); cell.font = INP; cell.fill = YEL
        cell.number_format = "0.0" if col == 11 else "0"
hb["A16"] = "Saisir les nombres de la semaine (pas les cumuls) ; « Pilotes actifs », « Note » et « Incidents » = valeur de la semaine."
hb["A16"].font = Font(name=F, italic=True)

# --- Prospects
pr = wb.create_sheet("Prospects")
header(pr, 1, [f"{a} / {b}" if a != b else a for a, b in zip(PROSPECT_FR, PROSPECT_EN)], [9, 12, 30, 12, 12, 16, 14, 15, 20, 14, 13, 13, 7, 15, 16, 16, 13, 14, 30])
dvs = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
pr.add_data_validation(dvs); dvs.add("N2:N2000")
pr.freeze_panes = "A2"

# --- Partenaires
pa = wb.create_sheet("Partenaires")
header(pa, 1, [f"{a} / {b}" if a != b else a for a, b in zip(PARTNER_FR, PARTNER_EN)], [8, 28, 14, 12, 13, 12, 8, 8, 8, 12, 14, 14, 14, 28])
pa.freeze_panes = "A2"

# --- Tableau de bord (dashboard), first sheet
db = wb.create_sheet("Tableau de bord", 0)
db["A1"] = "Nhill — tableau de bord commercial (90 jours) / sales dashboard"; db["A1"].font = Font(name=F, bold=True, size=14, color="173F5F")
db["A3"] = "Date de début (lundi S1) / Start date (Monday W1)"; db["A3"].font = B
db["B3"] = None; db["B3"].font = INP; db["B3"].fill = YEL; db["B3"].number_format = "yyyy-mm-dd"
db["B3"].comment = Comment("Saisir la date du lundi de la semaine 1 (AAAA-MM-JJ).", "Nhill")
header(db, 5, ["Indicateur / Indicator", "Réalisé / Actual", "Objectif / Target", "Atteinte / Progress", "Définition"], [44, 16, 14, 14, 60])
H_ = "Hebdo!"
kpis = [
    ("Prospects identifiés", f"=SUM({H_}C2:C14)", 150, "lignes ajoutées au suivi"),
    ("Contacts réalisés", f"=SUM({H_}D2:D14)", 90, "statut ≥ contacté"),
    ("Démos faites", f"=SUM({H_}E2:E14)", 25, "statut ≥ démo faite"),
    ("Taux contact → démo", f"=IFERROR(B8/B7,0)", 0.25, "démos / contacts"),
    ("Essais acceptés (pilotes)", f"=SUM({H_}F2:F14)", 10, "conditions de l'essai acceptées par écrit"),
    ("Taux démo → pilote", f"=IFERROR(B10/B8,0)", 0.40, "essais acceptés / démos"),
    ("Pilotes actifs (dernière semaine saisie)", f"=IFERROR(LOOKUP(2,1/({H_}G2:G14<>\"\"),{H_}G2:G14),0)", 10, "≥ 4 jours d'usage la semaine passée"),
    ("Clients", f"=SUM({H_}I2:I14)", 4, "premier paiement confirmé"),
    ("Taux pilote → client", f"=IFERROR(B13/SUM({H_}H2:H14),0)", 0.40, "clients / pilotes terminés"),
    ("Partenaires signés", f"=SUM({H_}J2:J14)", 3, "contrats signés"),
    ("Note de recommandation moyenne", f"=IFERROR(AVERAGE({H_}K2:K14),0)", 7, "moyenne 0-10"),
    ("Incidents bloquants > 48 h (cumul)", f"=SUM({H_}L2:L14)", 0, "doit rester à 0"),
    ("Dépensé (USD)", f"={SPENT_TOTAL}", f"={BUDGET_TOTAL}", "vs enveloppe du scénario (feuille Budget)"),
    ("Coût par démo (USD)", f"=IFERROR(B18/B8,0)", None, "dépenses / démos"),
]
pct_rows = []
for i, (name, f_, target, d) in enumerate(kpis):
    r = 6 + i
    db.cell(row=r, column=1, value=name).font = NORM
    db.cell(row=r, column=2, value=f_).font = NORM
    if target is not None:
        c = db.cell(row=r, column=3, value=target); c.font = INP if not str(target).startswith("=") else NORM
        if str(target).startswith("="): c.font = Font(name=F, color="008000")
    if target not in (None, 0):
        db.cell(row=r, column=4, value=f"=IFERROR(B{r}/C{r},0)").number_format = "0%"
    db.cell(row=r, column=5, value=d).font = NORM
    if "Taux" in name:
        for col in (2, 3): db.cell(row=r, column=col).number_format = "0.0%"
    elif "Note" in name:
        db.cell(row=r, column=2).number_format = "0.0"
    elif "USD" in name:
        for col in (2, 3): db.cell(row=r, column=col).number_format = "#,##0.00;(#,##0.00);-"
    else:
        db.cell(row=r, column=2).number_format = "0"
db["A21"] = "Objectifs : plan de 90 jours (marketing/fr/03-plan-90-jours.md). Ne jamais additionner USD et CDF."
db["A21"].font = Font(name=F, italic=True)
db["A22"] = "Mode d'emploi : saisir la date de début ci-dessus, puis chaque lundi la ligne de la semaine dans « Hebdo » ; les dépenses dans « Budget »."
db["A22"].font = Font(name=F, italic=True)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold if c.font else False, italic=c.font.italic if c.font else False,
                              color=c.font.color if c.font else None, size=c.font.size if c.font else 11)
wb.save(os.path.join(OUT, "kpi-tableau-de-bord.xlsx"))
print("ok")
